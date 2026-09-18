"""The PDF / Excel export follows the search on screen, and refuses politely.

A receptionist pressing the button used to be sent to the login page and, being
signed in already, on to the old dashboard — the button seemed to wander off.
"""

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse

from accounts.models import ClinicRole
from branches.models import Branch
from patients.models import Patient
from tenants.context import tenant_context
from tenants.models import Tenant

User = get_user_model()


class PatientExportTests(TestCase):
    def setUp(self):
        self.tenant = Tenant.objects.first()
        with tenant_context(self.tenant):
            self.branch = Branch.all_objects.create(tenant=self.tenant, name="Ex", code="EX")
            roles = {n: ClinicRole.all_objects.get_or_create(tenant=self.tenant, name=n)[0] for n in ("Admin", "Reception")}
            Patient.all_objects.create(tenant=self.tenant, name="Adam Export", gender="male", branch=self.branch)
            Patient.all_objects.create(tenant=self.tenant, name="Eve Export", gender="female", branch=self.branch)
        for key, role in (("admin", "Admin"), ("desk", "Reception")):
            User.objects.create_user(
                username=key, email=f"{key}@ex.local", password="pass12345",
                tenant=self.tenant, role=roles[role], branch=self.branch,
            )

    def login(self, key):
        self.client.logout()
        self.assertTrue(self.client.login(email=f"{key}@ex.local", password="pass12345"))

    def test_the_front_desk_gets_a_403_not_a_trip_through_the_login_page(self):
        self.login("desk")
        response = self.client.get(reverse("patients:patient_list_export"), {"export": "excel"})
        self.assertEqual(response.status_code, 403)

    def test_management_exports_only_what_the_search_matches(self):
        import io

        from openpyxl import load_workbook

        self.login("admin")

        def names(**params):
            response = self.client.get(reverse("patients:patient_list_export"), {"export": "excel", **params})
            self.assertEqual(response.status_code, 200)
            sheet = load_workbook(io.BytesIO(response.content)).active
            return {row[1] for row in sheet.iter_rows(min_row=1, values_only=True)} & {"Adam Export", "Eve Export"}

        self.assertEqual(names(), {"Adam Export", "Eve Export"})
        self.assertEqual(names(gender="female"), {"Eve Export"})
        self.assertEqual(names(search="adam"), {"Adam Export"})
