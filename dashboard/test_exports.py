"""Coverage for the two third-party libraries that had none: reportlab and pillow.

These live here rather than beside the code they test because `utils/` has no
`__init__.py` — it is an implicit namespace package, and test discovery through
one is unreliable. `dashboard` is an installed app and already hosts the other
cross-cutting checks (static-asset integrity).

Why they exist at all: `reportlab` and `pillow` are the only two dependencies
that produce binary output, and nothing verified either. That makes a version
bump unfalsifiable — the suite would stay green while PDF export produced
garbage, because no test ever looked at the bytes. Written before bumping
either, so the bump has something to fail against.

The Arabic rendering tests were added with FE-016; see
static/fonts/cairo/README.md for why the font is tracked and what had to be
verified about it.

Note for whoever touches the PDF path next: `utils/utils.py::export_pdf` is the
live one, used by billing, patients and services. `utils/export_pdf.py` defines
a *different* function with the same name and is imported by nothing — dead
code, left alone here rather than deleted as an unrelated change.
"""

import tempfile
from io import BytesIO

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, override_settings

from utils.utils import (
    ARABIC_FONT,
    EXCEL_SHEET_NAME_LIMIT,
    excel_sheet_name,
    ARABIC_FONT_PATH,
    arabic_table_style,
    export_excel,
    export_pdf,
    register_arabic_font,
    rtl,
)

HEADERS = ["الاسم", "الهاتف", "المبلغ"]
ROWS = [
    ["أحمد محمد", "01001234567", "1200.00"],
    ["منى علي", "01112345678", "450.50"],
]


class PdfExportTests(SimpleTestCase):
    """reportlab. The assertions are about the bytes, not just the status code —
    a broken PDF still comes back as a 200 with a PDF content type."""

    def response(self):
        return export_pdf(ROWS, HEADERS, "تقرير المرضى", "patients-report")

    def test_the_response_is_served_as_a_pdf_attachment(self):
        response = self.response()
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertIn("patients-report.pdf", response["Content-Disposition"])

    def test_the_body_is_a_structurally_valid_pdf(self):
        body = self.response().content
        self.assertTrue(body.startswith(b"%PDF-"), "missing the PDF magic header")
        self.assertIn(b"%%EOF", body, "PDF is not terminated")
        # A document containing a title and a 3x3 table is not 200 bytes. This
        # catches an "empty but well-formed" regression, which is what a
        # platypus API change would most likely produce.
        self.assertGreater(len(body), 1000, f"suspiciously small PDF: {len(body)} bytes")

    def test_an_empty_dataset_still_produces_a_document(self):
        """Reached whenever a report matches nothing — a crash here would be a
        500 on an ordinary empty search."""
        body = export_pdf([], HEADERS, "فارغ", "empty").content
        self.assertTrue(body.startswith(b"%PDF-"))


class ExcelExportTests(SimpleTestCase):
    """openpyxl. Read back through openpyxl rather than trusting the bytes, so
    the test fails if the sheet is written but empty."""

    def test_the_workbook_round_trips_with_its_data(self):
        response = export_excel(ROWS, HEADERS, "المرضى", "patients-report")
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn("patients-report.xlsx", response["Content-Disposition"])

        body = response.content
        self.assertTrue(body.startswith(b"PK"), "not a zip container, so not xlsx")

        workbook = openpyxl.load_workbook(BytesIO(body))
        sheet = workbook.active
        values = [[cell.value for cell in row] for row in sheet.iter_rows()]
        flat = [v for row in values for v in row if v is not None]

        self.assertIn("المرضى", flat)            # the title row
        self.assertIn("الاسم", flat)             # the header row
        self.assertIn("أحمد محمد", flat)         # a data row, Arabic intact
        self.assertIn("01001234567", flat)


class ImageFieldTests(TestCase):
    """pillow. Django uses it to validate and to read image dimensions, so a
    pillow bump can break uploads without any of our code changing."""

    def png_bytes(self):
        """Generated with pillow rather than hard-coded, so the test exercises
        the installed version at both ends."""
        from PIL import Image

        buffer = BytesIO()
        Image.new("RGB", (32, 24), (200, 30, 30)).save(buffer, format="PNG")
        return buffer.getvalue()

    def test_django_accepts_a_real_image(self):
        """This is the pillow integration point: forms.ImageField opens the
        upload with pillow and calls verify()."""
        from django import forms

        cleaned = forms.ImageField().clean(
            SimpleUploadedFile("logo.png", self.png_bytes(), content_type="image/png")
        )
        self.assertEqual(cleaned.image.size, (32, 24))
        self.assertEqual(cleaned.image.format, "PNG")

    def test_django_rejects_a_file_that_is_not_an_image(self):
        """The other half — a bump that silently stopped validating would be
        worse than one that broke loudly."""
        from django import forms

        with self.assertRaises(Exception):
            forms.ImageField().clean(
                SimpleUploadedFile("evil.png", b"not an image at all",
                                   content_type="image/png")
            )

    def test_a_patient_photo_saves_and_reads_back(self):
        """End to end through the model field and storage. MEDIA_ROOT is
        redirected so the suite does not litter the project's media/ directory."""
        from patients.models import Patient

        from tenants.context import tenant_context
        from tenants.models import Tenant

        tenant = Tenant.objects.first()
        with tempfile.TemporaryDirectory() as media_root:
            with override_settings(MEDIA_ROOT=media_root):
                with tenant_context(tenant):
                    patient = Patient.all_objects.create(
                        tenant=tenant, name="Photo Patient",
                        photo=SimpleUploadedFile(
                            "p.png", self.png_bytes(), content_type="image/png"
                        ),
                    )
                    patient.refresh_from_db()
                    self.assertTrue(patient.photo.name.endswith(".png"))
                    self.assertEqual(patient.photo.width, 32)
                    self.assertEqual(patient.photo.height, 24)
                    with patient.photo.open("rb") as handle:
                        self.assertTrue(handle.read().startswith(b"\x89PNG"))


# Real strings this application exports: report titles, table headers, seeded
# patient names, diagnoses, medications, allergens, branch names, plan titles,
# and the choice labels — plus money, phone numbers and serials, which must
# survive untouched.
ARABIC_CORPUS = [
    "قائمة المرضى - فرع سوهاج", "تقرير مالي", "قائمة المدفوعات", "قائمة الخدمات",
    "الاسم", "الهاتف", "المبلغ", "التاريخ", "رقم الإيصال", "الفرع", "الطبيب",
    "إجمالي الإيرادات", "المصروفات", "صافي الربح", "طريقة الدفع",
    "حكة وطفح جلدي", "تساقط الشعر", "بقع داكنة بالوجه", "حب الشباب", "جفاف الجلد",
    "التهاب جلدي تحسسي", "أكزيما", "حب شباب متوسط", "تصبغات جلدية", "صدفية خفيفة",
    "كريم هيدروكورتيزون", "لوراتادين", "دوكسيسيكلين", "مرطب طبي",
    "البنسلين", "الأسبرين", "اليود", "اللاتكس", "السلفا",
    "طفح جلدي", "تورم", "ضيق تنفس", "حكة شديدة",
    "أحمد محمد", "منى علي", "فاطمة عبد الرحمن", "محمود إبراهيم",
    "سوهاج", "أسيوط", "القاهرة", "المعادي",
    "علاج بالليزر", "جلسة ليزر Eximer 160", "كشف عيادة عام", "استشارة جلدية",
    "مرتين يومياً", "مرة يومياً", "عند اللزوم", "بعد الأكل", "قبل النوم",
    "خطة علاج", "جلسة علاج", "مكتملة", "مجدولة", "ملغاة", "لم يحضر",
    "ذكر", "أنثى", "أعزب", "متزوج", "مطلق", "أرمل",
]

NON_ARABIC = ["1200.00", "01001234567", "20260908-001", "Eximer 120", "", "0", "-50.25"]


class ArabicPdfRenderingTests(SimpleTestCase):
    """FE-016. Arabic in a PDF needs three separate things to be right — an
    embedded font with the glyphs, contextual shaping, and bidi reordering — and
    getting two of the three produces a document that looks fixed while being
    unreadable. Each is asserted here rather than assumed.
    """

    def font_face(self):
        register_arabic_font()
        from reportlab.pdfbase import pdfmetrics

        return pdfmetrics.getFont(ARABIC_FONT).face

    def test_the_font_is_present_and_registers(self):
        self.assertTrue(
            ARABIC_FONT_PATH.exists(),
            f"the Arabic font is missing from the repository: {ARABIC_FONT_PATH}",
        )
        from reportlab.pdfbase import pdfmetrics

        register_arabic_font()
        self.assertIn(ARABIC_FONT, pdfmetrics.getRegisteredFontNames())

    def test_registering_twice_is_harmless(self):
        """reportlab's font registry is process-global and these exports run per
        request, so this happens on every call after the first."""
        self.assertEqual(register_arabic_font(), ARABIC_FONT)
        self.assertEqual(register_arabic_font(), ARABIC_FONT)

    def test_the_font_actually_covers_arabic(self):
        """Guards against a replacement that is Arabic in name only. The font
        previously bundled as Cairo-Regular.ttf was a Latin-only subset."""
        face = self.font_face()
        for label, codepoint in (
            ("alef", 0x0627), ("beh", 0x0628), ("meem", 0x0645),
            ("heh", 0x0647), ("lam", 0x0644), ("hamza-alef", 0x0623),
        ):
            with self.subTest(letter=label):
                self.assertIsNotNone(
                    face.charToGlyph.get(codepoint),
                    f"no glyph for U+{codepoint:04X} ({label})",
                )

    def test_every_codepoint_the_shaper_emits_has_a_glyph(self):
        """The test that matters, and the one that failed first.

        "Does the font have Arabic" and "can it draw what the shaper produces"
        are different questions. arabic_reshaper maps letters onto Presentation
        Forms-B, and Cairo implements isolated forms through the base codepoint
        plus GSUB rather than duplicating them there — so 17 codepoints had no
        glyph until the reshaper was configured with
        use_unshaped_instead_of_isolated. A blank glyph is silent: the PDF is
        still valid, it just has nothing where the patient's name should be.
        """
        face = self.font_face()
        missing = {}
        for text in ARABIC_CORPUS:
            for char in rtl(text):
                if face.charToGlyph.get(ord(char)) is None:
                    missing.setdefault(f"U+{ord(char):04X} {char!r}", text)
        self.assertEqual(
            missing, {},
            "these codepoints would render blank:\n"
            + "\n".join(f"  {k} — e.g. in {v!r}" for k, v in sorted(missing.items())),
        )

    def test_shaping_and_reordering_actually_happen(self):
        """Without this, a no-op rtl() would satisfy the coverage test above."""
        original = "أحمد محمد"
        rendered = rtl(original)
        self.assertNotEqual(rendered, original, "text was not shaped at all")
        self.assertEqual(len(rendered), len(original), "characters were lost")
        # Reordered: the first character drawn is the last letter of the source.
        self.assertEqual(rendered[0], "ﺪ", "bidi reordering did not run")

    def test_non_arabic_values_pass_through_untouched(self):
        """Amounts, receipt numbers and serials must not be reordered."""
        for value in NON_ARABIC:
            with self.subTest(value=value):
                self.assertEqual(rtl(value), value)

    def test_none_and_numbers_are_handled(self):
        """Table cells arrive straight from querysets, so they are not all str."""
        self.assertEqual(rtl(None), "")
        self.assertEqual(rtl(0), "0")
        self.assertEqual(rtl(1200), "1200")

    def test_the_pdf_embeds_the_arabic_font(self):
        """End to end: the produced document must carry a Cairo subset, not just
        reference a font by name."""
        import re

        body = export_pdf(ROWS, HEADERS, "تقرير المرضى", "r").content
        fonts = {f.decode() for f in re.findall(rb"/BaseFont\s*/([A-Za-z0-9+,.-]+)", body)}
        self.assertTrue(
            any("Cairo" in name for name in fonts),
            f"no Cairo subset embedded; fonts present: {sorted(fonts)}",
        )

    def test_arabic_content_makes_the_document_substantially_larger(self):
        """A subset font embed is tens of kilobytes. If Arabic silently stopped
        being drawn, the document would collapse back to its old size."""
        body = export_pdf(ROWS, HEADERS, "تقرير المرضى", "r").content
        self.assertGreater(
            len(body), 8000,
            f"no font appears to be embedded: {len(body)} bytes",
        )

    def test_every_font_in_the_table_style_is_the_arabic_face(self):
        """Closes a gap the first version of these tests had: asserting the PDF
        embeds Cairo does not prove the *table* uses it. The title alone embeds
        the font, so a Helvetica-Bold header row — which is what this code
        originally had — passed the embed check while rendering every header cell
        blank. The style is where that decision lives, so it is checked here.
        """
        commands = arabic_table_style().getCommands()
        fontnames = [c for c in commands if c[0].upper() == "FONTNAME"]
        self.assertTrue(fontnames, "the table style sets no font at all")
        for command in fontnames:
            with self.subTest(command=command):
                self.assertEqual(
                    command[-1], ARABIC_FONT,
                    f"{command[-1]!r} cannot draw Arabic; only {ARABIC_FONT} can",
                )

    def test_the_table_style_covers_the_header_row(self):
        """The header contains Arabic too, and it was the row that regressed."""
        for command in arabic_table_style().getCommands():
            if command[0].upper() == "FONTNAME":
                start, stop = command[1], command[2]
                if start == (0, 0) and stop == (-1, -1):
                    return
        self.fail("no FONTNAME command spans the whole table including row 0")


class ExcelSheetNameTests(SimpleTestCase):
    """Excel rejects sheet names over 31 characters or containing : \ / ? * [ ].

    openpyxl only *warns*, so the file gets written and then fails to open in
    some readers — a defect that surfaces at the client, not in the logs. Found
    during the clean-clone handover rehearsal: the real title
    "قائمة المرضى - فرع الفرع الرئيسي" is 32 characters.
    """

    def test_a_long_arabic_title_is_shortened(self):
        title = "قائمة المرضى - فرع الفرع الرئيسي"
        self.assertGreater(len(title), EXCEL_SHEET_NAME_LIMIT)
        self.assertLessEqual(len(excel_sheet_name(title)), EXCEL_SHEET_NAME_LIMIT)

    def test_forbidden_characters_are_removed(self):
        for bad in ":\/?*[]":
            with self.subTest(char=bad):
                self.assertNotIn(bad, excel_sheet_name(f"a{bad}b"))

    def test_an_empty_title_still_yields_a_valid_name(self):
        self.assertTrue(excel_sheet_name(""))
        self.assertTrue(excel_sheet_name(None))

    def test_a_workbook_with_a_long_title_opens_cleanly(self):
        """The end-to-end claim: no warning, and openpyxl can read it back."""
        import warnings

        with warnings.catch_warnings():
            warnings.simplefilter("error")  # a UserWarning here fails the test
            response = export_excel(
                ROWS, HEADERS, "قائمة المرضى - فرع الفرع الرئيسي", "long-title"
            )
        workbook = openpyxl.load_workbook(BytesIO(response.content))
        self.assertLessEqual(len(workbook.active.title), EXCEL_SHEET_NAME_LIMIT)
        # The full title is not lost — it is still the first row.
        first_row = [c.value for c in next(workbook.active.iter_rows())]
        self.assertIn("قائمة المرضى - فرع الفرع الرئيسي", first_row)
